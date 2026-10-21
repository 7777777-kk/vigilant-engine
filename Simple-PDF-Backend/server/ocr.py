"""Bounded, offline Chinese/English OCR with editable text and ruled tables."""
from __future__ import annotations

import importlib.util
import math
import time
from copy import deepcopy
from io import BytesIO
from pathlib import Path

MAX_OCR_PAGES = 10
MAX_OCR_PIXELS = 12_000_000
OCR_SECONDS = 145


def capability() -> dict:
    spec = importlib.util.find_spec('rapidocr_onnxruntime')
    models = Path(spec.origin).parent / 'models' if spec and spec.origin else None
    available = bool(models and len(list(models.glob('*.onnx'))) >= 3
                     and importlib.util.find_spec('onnxruntime'))
    return {'available': available, 'engine': 'RapidOCR', 'languages': ['zh-Hans', 'en'],
            'maxPages': MAX_OCR_PAGES, 'reason': '' if available else '云端尚未安装 OCR 引擎，请部署 OCR 更新包。'}


def make_engine():
    if not capability()['available']:
        raise ValueError('此页需要 OCR，云端尚未安装 OCR 引擎，请部署 OCR 更新包。')
    try:
        from rapidocr_onnxruntime import RapidOCR
        return RapidOCR(intra_op_num_threads=1, inter_op_num_threads=1,
                        det_limit_side_len=1600, text_score=0.35)
    except Exception as exc:
        raise ValueError('OCR 引擎启动失败，请检查服务器模型与运行依赖。') from exc


def render_page(path: Path, index: int):
    import pypdfium2 as pdfium
    with pdfium.PdfDocument(path) as document:
        page = document[index]
        try:
            width, height = page.get_size()
            if not all(math.isfinite(v) and v > 0 for v in (width, height)):
                raise ValueError('PDF 页面尺寸无效。')
            scale = min(200 / 72, math.sqrt(MAX_OCR_PIXELS / (width * height)))
            bitmap = page.render(scale=scale)
            try:
                return bitmap.to_pil().convert('RGB')
            finally:
                bitmap.close()
        finally:
            page.close()


def _positions(values) -> list[int]:
    groups = []
    for value in values:
        value = int(value)
        if not groups or value - groups[-1][-1] > 4:
            groups.append([value])
        else:
            groups[-1].append(value)
    return [round(sum(group) / len(group)) for group in groups]


def ruled_tables(image) -> list[dict]:
    """Find simple axis-aligned grids. Complex/merged cells remain approximate."""
    import cv2
    import numpy as np
    gray = cv2.cvtColor(np.asarray(image), cv2.COLOR_RGB2GRAY)
    mask = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_MEAN_C, cv2.THRESH_BINARY_INV, 31, 15)
    height, width = gray.shape
    horizontal = cv2.morphologyEx(mask, cv2.MORPH_OPEN,
                                cv2.getStructuringElement(cv2.MORPH_RECT, (max(30, width // 35), 1)))
    vertical = cv2.morphologyEx(mask, cv2.MORPH_OPEN,
                              cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(30, height // 35))))
    grid = cv2.bitwise_or(horizontal, vertical)
    contours, _ = cv2.findContours(grid, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    tables = []
    for contour in contours:
        x, y, w, h = cv2.boundingRect(contour)
        if w < width * .15 or h < 35:
            continue
        xs = _positions(np.flatnonzero(np.count_nonzero(vertical[y:y+h, x:x+w], axis=0) >= h * .5))
        ys = _positions(np.flatnonzero(np.count_nonzero(horizontal[y:y+h, x:x+w], axis=1) >= w * .5))
        if not (2 <= len(xs) <= 21 and 2 <= len(ys) <= 101 and (len(xs)-1)*(len(ys)-1) <= 600):
            continue
        if any(b-a < 12 for a, b in zip(xs, xs[1:])) or any(b-a < 12 for a, b in zip(ys, ys[1:])):
            continue
        tables.append({'x': [x + value for value in xs], 'y': [y + value for value in ys]})
    return sorted(tables, key=lambda table: table['y'][0])[:10]


def recognize(engine, image) -> list[dict]:
    import numpy as np
    try:
        result, _ = engine(np.asarray(image))
    except Exception as exc:
        raise ValueError('OCR 无法识别此页，请使用更清晰、方向正确的扫描件。') from exc
    words = []
    for box, text, score in result or []:
        text = ''.join(char for char in str(text) if char in '\n\t' or ord(char) >= 32).strip()
        if not text or float(score) < .35:
            continue
        xs, ys = [float(p[0]) for p in box], [float(p[1]) for p in box]
        words.append({'text': text, 'score': float(score), 'x': min(xs), 'y': min(ys),
                      'right': max(xs), 'bottom': max(ys)})
    return sorted(words, key=lambda word: (word['y'], word['x']))


def lines(words: list[dict]) -> list[str]:
    rows = []
    for word in sorted(words, key=lambda item: (item['y'], item['x'])):
        center = (word['y'] + word['bottom']) / 2
        row = next((row for row in reversed(rows[-3:]) if abs(row[0] - center) < max(5, (word['bottom']-word['y']) * .55)), None)
        if row is None:
            rows.append([center, [word]])
        else:
            row[1].append(word)
    return [' '.join(word['text'] for word in sorted(row[1], key=lambda item: item['x'])) for row in rows]


def append_ocr_page(document, words: list[dict], tables: list[dict]):
    assigned = set()
    blocks = []
    for bounds in tables:
        xs, ys = bounds['x'], bounds['y']
        cells = [[[] for _ in range(len(xs)-1)] for _ in range(len(ys)-1)]
        for index, word in enumerate(words):
            cx, cy = (word['x']+word['right'])/2, (word['y']+word['bottom'])/2
            col = next((i for i in range(len(xs)-1) if xs[i] <= cx <= xs[i+1]), None)
            row = next((i for i in range(len(ys)-1) if ys[i] <= cy <= ys[i+1]), None)
            if row is not None and col is not None and index not in assigned:
                cells[row][col].append(word)
                assigned.add(index)
        blocks.append((ys[0], 'table', cells))
    for index, word in enumerate(words):
        if index not in assigned:
            blocks.append((word['y'], 'text', word))
    pending = []
    for _, kind, value in sorted(blocks, key=lambda block: block[0]):
        if kind == 'text':
            pending.append(value)
            continue
        for text in lines(pending):
            document.add_paragraph(text)
        pending = []
        table = document.add_table(rows=len(value), cols=len(value[0]))
        table.style = 'Table Grid'
        for row, cells in enumerate(value):
            for col, cell in enumerate(cells):
                table.cell(row, col).text = '\n'.join(lines(cell))
    for text in lines(pending):
        document.add_paragraph(text)


def append_native_page(target, source):
    """Copy pdf2docx page content while preserving images and table styles."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml.ns import qn
    for style in source.styles.element:
        identifier = style.get(qn('w:styleId'))
        if identifier and not any(s.get(qn('w:styleId')) == identifier for s in target.styles.element):
            target.styles.element.append(deepcopy(style))
    mapping = {}
    for relation in source.part.rels.values():
        if relation.reltype == RT.IMAGE:
            mapping[relation.rId] = target.part.get_or_add_image(BytesIO(relation.target_part.blob))[0]
        elif relation.is_external:
            mapping[relation.rId] = target.part.relate_to(relation.target_ref, relation.reltype, is_external=True)
    for element in source.element.body:
        if element.tag == qn('w:sectPr'):
            continue
        copied = deepcopy(element)
        for node in copied.iter():
            for key, value in list(node.attrib.items()):
                if key.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}') and value in mapping:
                    node.set(key, mapping[value])
        target.element.body.insert_element_before(copied, 'w:sectPr')


def process_pages(path: Path, indices: list[int], text_pages: list[str], output: Path, as_word: bool) -> str:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt
    empty = [i for i, text in zip(indices, text_pages) if not text.strip()]
    if len(empty) > MAX_OCR_PAGES:
        raise ValueError(f'一次最多识别 {MAX_OCR_PAGES} 页扫描件，请在页码范围中分批转换。')
    engine = make_engine()
    deadline = time.monotonic() + OCR_SECONDS
    document = Document() if as_word else None
    if document is not None:
        style = document.styles['Normal']
        style.font.name = 'Calibri'
        style.font.size = Pt(11)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    all_text, unrecognized, low_confidence = [], [], False
    with __import__('tempfile').TemporaryDirectory(prefix='ocr-', dir=output.parent) as temporary:
        for sequence, (index, native_text) in enumerate(zip(indices, text_pages)):
            if time.monotonic() > deadline:
                raise ValueError('OCR 处理时间较长，请减少页数后分批识别。')
            if document is not None and sequence:
                document.add_page_break()
            if native_text.strip():
                all_text.append(native_text)
                if document is not None:
                    from pdf2docx import Converter
                    native_output = Path(temporary) / 'native.docx'
                    converter = Converter(str(path))
                    try:
                        converter.convert(str(native_output), pages=[index], multi_processing=False)
                    finally:
                        converter.close()
                    append_native_page(document, Document(native_output))
                continue
            image = render_page(path, index)
            try:
                words = recognize(engine, image)
                tables = ruled_tables(image) if document is not None else []
            finally:
                image.close()
            if not words:
                unrecognized.append(index + 1)
                text = f'[第 {index + 1} 页未识别出文字，请核对原页]'
                all_text.append(text)
                if document is not None:
                    document.add_paragraph(text)
                continue
            low_confidence |= any(word['score'] < .75 for word in words)
            all_text.append('\n'.join(lines(words)))
            if document is not None:
                append_ocr_page(document, words, tables)
    if len(unrecognized) == len(indices):
        raise ValueError('OCR 未识别出文字，请使用更清晰、方向正确的扫描件；手写内容可能无法识别。')
    if document is not None:
        document.save(output)
    else:
        output.write_text('\n\n'.join(all_text), encoding='utf-8')
    warning = f'已自动 OCR 识别 {len(empty)} 页扫描件。请校对姓名、数字和表格；复杂排版、合并单元格及手写内容可能有误。'
    if low_confidence:
        warning += '部分内容识别置信度较低，请重点校对。'
    if unrecognized:
        warning += '第 ' + '、'.join(map(str, unrecognized)) + ' 页未识别出文字，已在结果中标注。'
    return warning
