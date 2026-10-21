"""Real PDF conversions, using upstream engines rather than extension changes."""
from __future__ import annotations

import io
import math
import os
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

from PIL import Image, ImageOps
from pypdf import PdfReader, PdfWriter

MAX_PAGES = 100
MAX_IMAGE_PIXELS = 24_000_000
MAX_OUTPUT_BYTES = 80 * 1024 * 1024
Image.MAX_IMAGE_PIXELS = MAX_IMAGE_PIXELS


def ocr_text_pages(path: Path, reader, indices: list[int], force: bool = False) -> list[str]:
    """Detect image-heavy scans even when a small header/logo has native text."""
    import pymupdf
    result = []
    with pymupdf.open(path) as document:
        for index in indices:
            page = document[index]
            text = (reader.pages[index].extract_text() or '').strip()
            meaningful = sum(char.isalnum() for char in page.get_text())
            image_heavy = False
            if meaningful < 200:
                for image in page.get_images():
                    if any(rect.get_area() > page.rect.get_area() * .5 for rect in page.get_image_rects(image[0])):
                        image_heavy = True
                        break
            result.append('' if force or image_heavy or meaningful == 0 else text)
    return result

TOOLS = {
    "pdf-word": {"title": "PDF 转 Word", "accept": ["pdf"], "min": 1, "max": 1},
    "word-pdf": {"title": "Word 转 PDF", "accept": ["docx"], "min": 1, "max": 1},
    "pdf-image": {"title": "PDF 转图片", "accept": ["pdf"], "min": 1, "max": 1},
    "image-pdf": {"title": "图片转 PDF", "accept": ["jpg", "jpeg", "png"], "min": 1, "max": 12},
    "merge": {"title": "合并 PDF", "accept": ["pdf"], "min": 2, "max": 8},
    "split": {"title": "拆分 / 提取页面", "accept": ["pdf"], "min": 1, "max": 1},
    "compress": {"title": "压缩 PDF", "accept": ["pdf"], "min": 1, "max": 1},
    "pdf-text": {"title": "PDF 转文字", "accept": ["pdf"], "min": 1, "max": 1},
}


def libreoffice_path() -> str | None:
    explicit = os.getenv("LIBREOFFICE_PATH", "")
    candidates = [explicit, shutil.which("libreoffice"), shutil.which("soffice"),
                  r"C:\Program Files\LibreOffice\program\soffice.exe"]
    return next((str(p) for p in candidates if p and Path(p).is_file()), None)


def read_pdf(path: Path) -> PdfReader:
    if path.read_bytes()[:5] != b"%PDF-":
        raise ValueError("文件内容不是有效的 PDF，请重新选择。")
    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("暂不支持加密 PDF，请先用原密码解密后再上传。")
        if not 1 <= len(reader.pages) <= MAX_PAGES:
            raise ValueError(f"PDF 页数须为 1–{MAX_PAGES} 页。")
        return reader
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("PDF 已损坏或无法解析，请更换文件。") from exc


def inspect_file(path: Path, extension: str) -> dict:
    if extension == "pdf":
        reader = read_pdf(path)
        # Text extraction is bounded here; actual conversions inspect all selected pages.
        has_text = any((p.extract_text() or "").strip() for p in list(reader.pages)[:3])
        return {"pages": len(reader.pages), "hasText": has_text}
    if extension == "docx":
        try:
            with zipfile.ZipFile(path) as archive:
                total = sum(entry.file_size for entry in archive.infolist())
                if total > MAX_OUTPUT_BYTES or len(archive.infolist()) > 5000:
                    raise ValueError("Word 文件解压后过大，请精简内容后重试。")
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("文件内容不是有效的 DOCX。")
            from docx import Document
            Document(path)
            return {"pages": None, "hasText": True}
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("Word 文件已损坏，仅支持 .docx 格式。") from exc
    if extension in ("png", "jpg", "jpeg"):
        try:
            with Image.open(path) as image:
                expected = "PNG" if extension == "png" else "JPEG"
                if image.format != expected:
                    raise ValueError("图片内容与扩展名不一致，请重新选择真实的 JPG 或 PNG。")
                if image.width * image.height > MAX_IMAGE_PIXELS:
                    raise ValueError("图片像素过大，请先缩小至 2400 万像素以内。")
                image.verify()
            return {"pages": 1, "hasText": False}
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("图片已损坏或尺寸过大。") from exc
    raise ValueError("仅支持 PDF、DOCX、JPG 和 PNG 文件。")


def page_indices(expression: str, count: int) -> list[int]:
    expression = expression.strip()
    if not expression:
        return list(range(count))
    if len(expression) > 600:
        raise ValueError("页码输入过长。")
    expression = expression.replace("，", ",").replace("–", "-")
    result: list[int] = []
    for part in expression.split(","):
        part = part.strip()
        if not re.fullmatch(r"[1-9]\d*(?:\s*-\s*[1-9]\d*)?", part):
            raise ValueError("页码格式错误，请输入 1-3,5 这样的范围。")
        bounds = [int(x.strip()) for x in part.split("-")]
        start, end = bounds[0], bounds[-1]
        if not 1 <= start <= end <= count:
            raise ValueError(f"页码超出范围或顺序错误，此文件共 {count} 页。")
        for index in range(start - 1, end):
            if index not in result:
                result.append(index)
    return result


def strip_active_content(writer: PdfWriter) -> None:
    """Do not propagate PDF JavaScript or launch actions into generated PDFs."""
    from pypdf.generic import NameObject
    writer.root_object.pop(NameObject("/OpenAction"), None)
    writer.root_object.pop(NameObject("/AA"), None)
    names = writer.root_object.get("/Names")
    if names:
        names.get_object().pop(NameObject("/JavaScript"), None)
        names.get_object().pop(NameObject("/EmbeddedFiles"), None)
    for page in writer.pages:
        page.pop(NameObject("/AA"), None)
        for annotation in page.get("/Annots", []):
            item = annotation.get_object()
            item.pop(NameObject("/AA"), None)
            action = item.get("/A")
            if action and action.get_object().get("/S") in ("/JavaScript", "/Launch"):
                item.pop(NameObject("/A"), None)


def render_image(pdf_path: Path, index: int, output: Path, dpi: int, image_format: str) -> None:
    import pypdfium2 as pdfium
    with pdfium.PdfDocument(pdf_path) as document:
        page = document[index]
        width, height = page.get_size()
        if not all(math.isfinite(value) and value > 0 for value in (width, height)):
            raise ValueError("PDF 页面尺寸无效。")
        scale = dpi / 72
        if width * height * scale * scale > MAX_IMAGE_PIXELS:
            raise ValueError("该页渲染后像素过大，请降低导出清晰度。")
        bitmap = page.render(scale=scale)
        try:
            image = bitmap.to_pil().convert("RGB")
            try:
                image.save(output, format="JPEG" if image_format == "jpg" else "PNG", quality=90)
            finally:
                image.close()
        finally:
            bitmap.close()
            page.close()


def _thumbnail(path: Path, output: Path) -> None:
    # Cap thumbnail output by page dimensions rather than using a fixed scale.
    import pypdfium2 as pdfium
    with pdfium.PdfDocument(path) as document:
        page = document[0]
        width, height = page.get_size()
        if not all(math.isfinite(v) and v > 0 for v in (width, height)):
            page.close()
            return
        bitmap = page.render(scale=min(480 / width, 640 / height, 1))
        try:
            image = bitmap.to_pil().convert("RGB")
            try:
                image.save(output, "PNG")
            finally:
                image.close()
        finally:
            bitmap.close()
            page.close()


def _write_manifest(output_dir: Path, files: list[tuple[Path, str]], warning: str, pages: int) -> dict:
    if sum(path.stat().st_size for path, _ in files) > MAX_OUTPUT_BYTES:
        raise ValueError("转换结果超过 80 MB，请减少页面或降低清晰度。")
    manifest = []
    for index, (path, name) in enumerate(files):
        item = {"id": str(index), "name": name, "disk": path.name,
                "size": path.stat().st_size, "format": path.suffix[1:]}
        manifest.append(item)
    pdf = next((path for path, _ in files if path.suffix == ".pdf"), None)
    if pdf:
        try:
            _thumbnail(pdf, output_dir / "preview.png")
        except Exception:
            pass  # A native PDF viewer remains available if the optional thumbnail fails.
    return {"files": manifest, "warning": warning, "pages": pages,
            "hasPreview": (output_dir / "preview.png").exists()}


def convert(tool: str, inputs: list[dict], options: dict, output_dir: str) -> dict:
    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    paths = [Path(item["path"]) for item in inputs]
    stem = Path(inputs[0]["name"]).stem[:60] or "document"
    files: list[tuple[Path, str]] = []
    warning = ""
    count = 0
    if tool == "image-pdf":
        # One image per PDF page, in the user-selected order; transparency is flattened to white.
        images = []
        try:
            for path in paths:
                with Image.open(path) as raw:
                    oriented = ImageOps.exif_transpose(raw)
                    rgba = oriented.convert("RGBA")
                    white = Image.new("RGB", rgba.size, "white")
                    white.paste(rgba, mask=rgba.getchannel("A"))
                    rgba.close()
                    images.append(white)
            output = directory / "result.pdf"
            images[0].save(output, "PDF", resolution=150.0, save_all=True, append_images=images[1:])
            files = [(output, f"{stem}-图片合并.pdf")]
            count = len(images)
        finally:
            for image in images:
                image.close()
    elif tool == "word-pdf":
        binary = libreoffice_path()
        if not binary:
            raise ValueError("服务器尚未安装 LibreOffice，Word 转 PDF 暂不可用。")
        profile = (directory / "lo-profile").resolve().as_uri()
        result = subprocess.run([binary, f"-env:UserInstallation={profile}", "--headless",
                                 "--convert-to", "pdf:writer_pdf_Export", "--outdir", str(directory),
                                 str(paths[0])], capture_output=True, timeout=120)
        expected = directory / (paths[0].stem + ".pdf")
        if result.returncode or not expected.exists():
            raise ValueError("Word 转换失败，请检查文档是否可以正常打开。")
        count = len(read_pdf(expected).pages)
        files = [(expected, f"{stem}.pdf")]
        warning = "字体替换可能影响换行，请预览转换结果。"
    else:
        reader = read_pdf(paths[0])
        indices = page_indices(options.get("pages", ""), len(reader.pages))
        count = len(indices)
        if tool == "pdf-word":
            text_pages = ocr_text_pages(paths[0], reader, indices, options.get('ocrMode') == 'force')
            if any(not text for text in text_pages):
                from .ocr import process_pages
                output = directory / 'result.docx'
                warning = process_pages(paths[0], indices, text_pages, output, as_word=True)
                return _write_manifest(directory, [(output, f'{stem}.docx')], warning, count)
            from pdf2docx import Converter
            output = directory / "result.docx"
            converter = Converter(str(paths[0]))
            try:
                converter.convert(str(output), pages=indices, multi_processing=False)
            finally:
                converter.close()
            from docx import Document
            doc = Document(output)
            if not any(p.text.strip() for p in doc.paragraphs) and not any(
                cell.text.strip() for table in doc.tables for row in table.rows for cell in row.cells
            ):
                from .ocr import process_pages
                warning = process_pages(paths[0], indices, [''] * len(indices), output, as_word=True)
                return _write_manifest(directory, [(output, f'{stem}.docx')], warning, count)
            warning = "复杂排版可能变化，请检查表格、分栏和字体。"
            if any(not page for page in text_pages):
                warning += "部分页面无可提取文字，可能包含扫描图片。"
            files = [(output, f"{stem}.docx")]
        elif tool == "pdf-image":
            fmt = options.get("imageFormat", "png")
            dpi = options.get("dpi", 144)
            if fmt not in ("png", "jpg") or dpi not in (96, 144, 200):
                raise ValueError("图片格式或清晰度设置无效。")
            if len(indices) > 30:
                raise ValueError("一次最多导出 30 页图片，请缩小页码范围。")
            total_size = 0
            for sequence, index in enumerate(indices):
                output = directory / f"page-{sequence + 1:03d}.{fmt}"
                render_image(paths[0], index, output, dpi, fmt)
                total_size += output.stat().st_size
                if total_size > MAX_OUTPUT_BYTES // 2:
                    raise ValueError("图片结果过大，请减少页面或降低清晰度。")
                files.append((output, f"{stem}-第{index + 1}页.{fmt}"))
            if len(files) > 1:
                archive_path = directory / "pages.zip"
                with zipfile.ZipFile(archive_path, "w", zipfile.ZIP_DEFLATED) as archive:
                    for path, name in files:
                        archive.write(path, arcname=name)
                files.append((archive_path, f"{stem}-图片合集.zip"))
        elif tool == "pdf-text":
            parts = ocr_text_pages(paths[0], reader, indices, options.get('ocrMode') == 'force')
            output = directory / "result.txt"
            if any(not part.strip() for part in parts):
                from .ocr import process_pages
                warning = process_pages(paths[0], indices, parts, output, as_word=False)
                return _write_manifest(directory, [(output, f'{stem}.txt')], warning, count)
            output.write_text("\n\n".join(parts), encoding="utf-8")
            files = [(output, f"{stem}.txt")]
            if any(not part.strip() for part in parts):
                warning = "部分页面无可提取文字，扫描图片中的文字不会出现在结果中。"
        elif tool in ("merge", "split", "compress"):
            writer = PdfWriter()
            if tool == "merge":
                count = 0
                for path in paths:
                    other = read_pdf(path)
                    count += len(other.pages)
                    if count > MAX_PAGES:
                        raise ValueError(f"合并后不能超过 {MAX_PAGES} 页。")
                    writer.append(other)
            else:
                if tool == "compress":
                    writer.clone_document_from_reader(reader)
                    count = len(reader.pages)
                    for page in writer.pages:
                        page.compress_content_streams()
                    writer.compress_identical_objects(remove_duplicates=True, remove_unreferenced=True)
                else:
                    for index in indices:
                        writer.add_page(reader.pages[index])
            strip_active_content(writer)
            output = directory / "result.pdf"
            with output.open("wb") as stream:
                writer.write(stream)
            suffix = {"merge": "合并", "split": "提取", "compress": "压缩"}[tool]
            files = [(output, f"{stem}-{suffix}.pdf")]
            if tool == "compress":
                if output.stat().st_size >= paths[0].stat().st_size:
                    warning = "原文件已较紧凑；无损优化未减小体积。已保留处理后的文件，不宣称压缩成功。"
                else:
                    saved = 1 - output.stat().st_size / paths[0].stat().st_size
                    warning = f"无损优化完成，文件体积减少 {saved:.1%}。"
        else:
            raise ValueError("不支持的转换功能。")
    return _write_manifest(directory, files, warning, count)


def worker_entry(send_connection, tool: str, inputs: list[dict], options: dict, output_dir: str) -> None:
    """Run native libraries in an isolated, timeout-controlled process."""
    try:
        send_connection.send({"ok": True, "result": convert(tool, inputs, options, output_dir)})
    except Exception as exc:
        message = str(exc) if isinstance(exc, ValueError) else "转换引擎无法处理此文件，请更换文件重试。"
        send_connection.send({"ok": False, "error": message})
    finally:
        send_connection.close()
